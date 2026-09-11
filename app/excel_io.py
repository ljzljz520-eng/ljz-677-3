"""Excel 读取、模板下载、异常导出。"""
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import catalog
from . import statuses as st

# 表头别名 -> 标准字段
HEADER_ALIASES = {
    "人员编号": "person_id", "人员id": "person_id", "人员ID": "person_id",
    "编号": "person_id", "体检编号": "person_id", "工号": "person_id",
    "身份证号": "person_id", "证件号": "person_id", "员工编号": "person_id",
    "项目": "item", "体检项目": "item", "项目名称": "item",
    "项目代码": "item", "检查项目": "item",
    "结果": "result", "体检结果": "result", "检查结果": "result",
    "结果值": "result", "数值": "result",
    "单位": "unit", "结果单位": "unit", "计量单位": "unit",
    "体检日期": "exam_date", "检查日期": "exam_date", "日期": "exam_date",
    "检测日期": "exam_date",
}
REQUIRED_FIELDS = ("person_id", "item", "result", "unit", "exam_date")
FIELD_TITLES = {
    "person_id": "人员编号", "item": "项目", "result": "结果",
    "unit": "单位", "exam_date": "体检日期",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ERR_FILL = PatternFill("solid", fgColor="FCE4E4")
WARN_FILL = PatternFill("solid", fgColor="FFF6D6")
DUP_FILL = PatternFill("solid", fgColor="EDE7F6")
TITLE_FONT = Font(bold=True, size=12)


class ExcelFormatError(Exception):
    """上传文件不是可解析的体检结果 Excel。"""


def _norm_header(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    for ch in " \t　*:：()（）":
        text = text.replace(ch, "")
    return text


def read_excel_rows(path: str):
    """读取 Excel，返回 (字段顺序, 原始行列表[{row_no, person_id,...}])。"""
    p = Path(path)
    if p.suffix.lower() not in {".xlsx"}:
        raise ExcelFormatError("仅支持 .xlsx 格式（Excel 2007 及以上），请另存后重试")
    if not p.exists() or p.stat().st_size == 0:
        raise ExcelFormatError("文件为空，请检查后重新上传")

    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对坏文件会抛多种异常
        raise ExcelFormatError(f"无法解析 Excel 文件：{exc}") from exc

    ws = wb.worksheets[0]
    row_iter = ws.iter_rows(values_only=True)

    mapping = None
    header_row_no = 0
    scanned = []
    for idx, raw_row in enumerate(row_iter, start=1):
        values = list(raw_row) if raw_row else []
        scanned.append((idx, values))
        found = {}
        for col_idx, cell in enumerate(values):
            key = _norm_header(cell)
            if key in HEADER_ALIASES:
                found[HEADER_ALIASES[key]] = col_idx
        if all(f in found for f in REQUIRED_FIELDS):
            mapping = found
            header_row_no = idx
            break
        if idx >= 5:
            break

    if mapping is None:
        missing_cols = "、".join(FIELD_TITLES[f] for f in REQUIRED_FIELDS)
        raise ExcelFormatError(
            f"未找到完整表头，必须包含：{missing_cols}（前 5 行内识别）。可先下载导入模板"
        )

    order = [f for f in REQUIRED_FIELDS]
    # 生成器此前已消费到表头行，继续读取剩余全部行
    for raw_row in row_iter:
        scanned.append((len(scanned) + 1, list(raw_row) if raw_row else []))

    rows = []
    # scanned 下标从 0 开始，header_row_no 为 1 基行号，切片正好跳过表头行
    for idx, values in scanned[header_row_no:]:
        if values is None:
            continue
        if all(v is None or str(v).strip() == "" for v in values):
            continue  # 整行空行跳过
        record = {"row_no": idx}
        for f in order:
            col = mapping[f]
            record[f] = values[col] if col < len(values) else None
        rows.append(record)

    wb.close()
    if not rows:
        raise ExcelFormatError("表头之后没有任何数据行")
    return order, rows


def _autosize(ws, widths=None):
    widths = widths or {}
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        width = widths.get(letter)
        if width is None:
            width = min(max((len(str(c.value)) for c in col_cells if c.value is not None), default=8) + 4, 48)
        ws.column_dimensions[letter].width = width


def build_template_workbook() -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "体检结果"
    headers = ["人员编号", "项目", "结果", "单位", "体检日期"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    samples = [
        ["P1001", "GLU", 5.2, "mmol/L", "2026-09-01"],
        ["P1001", "SBP", 128, "mmHg", "2026-09-01"],
        ["P1002", "WBC", 6.8, "10^9/L", "2026-09-01"],
    ]
    for row in samples:
        ws.append(row)

    note = wb.create_sheet("填写说明")
    notes = [
        ["填写说明"],
        ["1. 五个必填列：人员编号、项目、结果、单位、体检日期，表头需与模板一致。"],
        ["2. 项目可填写项目代码（如 GLU）或中文名称（如 空腹血糖）。"],
        ["3. 结果必须为数值；超出参考范围会提示但仍上送，超出物理可行区间则校验失败。"],
        ["4. 日期支持 2026-09-01、2026/09/01、20260901 三种写法。"],
        ["5. 同一人员、同一项目、同一天仅允许一条有效结果，重复行将被拦截并可在异常表导出。"],
        ["6. 单位必须与“项目目录”页中该项目的单位一致。"],
    ]
    for line in notes:
        note.append(line)
    note["A1"].font = TITLE_FONT
    note.column_dimensions["A"].width = 100

    cat = wb.create_sheet("项目目录")
    cat.append(["项目代码", "项目名称", "标准单位", "参考范围"])
    for col in range(1, 5):
        c = cat.cell(row=1, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for item in catalog.catalog_rows():
        cat.append([item["code"], item["name"], item["unit"], item["ref_range"]])

    _autosize(ws, {"A": 14, "B": 14, "C": 10, "D": 12, "E": 14})
    _autosize(cat)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# 异常导出每个页签对应的行状态
EXPORT_SHEETS = [
    ("校验失败", st.INVALID, ERR_FILL),
    ("平台异常", st.PLATFORM_ERROR, ERR_FILL),
    ("重复数据", st.DUPLICATE, DUP_FILL),
]


def build_error_workbook(task: dict, rows: list[dict]) -> BytesIO:
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "汇总"
    summary.append(["体检结果批量上送异常清单"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(["任务编号", task["id"]])
    summary.append(["源文件", task["filename"]])
    summary.append(["任务状态", st.TASK_STATUS.get(task["status"], task["status"])])
    summary.append(["数据总行数", task["total_count"]])
    summary.append(["上送成功", task["success_count"]])
    summary.append(["校验失败", task["invalid_count"]])
    summary.append(["平台异常", task["platform_error_count"]])
    summary.append(["重复数据", task["duplicate_count"]])
    for col in "AB":
        summary.column_dimensions[col].width = 22

    base_headers = ["Excel行号", "人员编号", "项目代码", "项目名称", "结果", "单位", "体检日期"]
    for sheet_name, status, fill in EXPORT_SHEETS:
        ws = wb.create_sheet(sheet_name)
        if status == st.INVALID:
            headers = base_headers + ["失败原因", "提示"]
        elif status == st.PLATFORM_ERROR:
            headers = base_headers + ["平台错误码", "平台返回信息"]
        else:
            headers = base_headers + ["重复原因"]
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
        part = [r for r in rows if r["status"] == status]
        for r in part:
            if status == st.INVALID:
                extra = [r["error_message"], r["warning"]]
            elif status == st.PLATFORM_ERROR:
                extra = [r["error_code"], r["error_message"]]
            else:
                extra = [r["error_message"]]
            ws.append(
                [r["row_no"], r["person_id"], r["item_code"], r["item_name"],
                 r["result"], r["unit"], r["exam_date"], *extra]
            )
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = fill
        ws.append([])
        ws.append([f"共 {len(part)} 行"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        _autosize(ws)
        ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
