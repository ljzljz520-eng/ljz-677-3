"""Excel 解析与字段标准化。"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook

from .catalog import canonical_unit

REQUIRED_HEADERS = ["人员编号", "项目", "结果", "单位", "体检日期"]
MAX_ROWS = 5000


class ExcelError(Exception):
    """文件层面的错误（无法解析、缺列等）。"""


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_excel(content: bytes) -> list[dict]:
    """读取第一个工作表，返回原始行字典列表（保留 row_no）。"""
    try:
        wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # openpyxl 抛错类型较多
        raise ExcelError(f"文件无法解析，请确认是有效的 .xlsx 文件（{e.__class__.__name__}）")

    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ExcelError("工作表为空")

    headers = [_cell_text(c) for c in header_row]
    index: dict[str, int] = {}
    for h in REQUIRED_HEADERS:
        if h not in headers:
            raise ExcelError(f"缺少必需列：{h}（应有列：{ '、'.join(REQUIRED_HEADERS) }）")
        index[h] = headers.index(h)

    records: list[dict] = []
    for i, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue  # 跳过完全空行

        def get(col: str) -> str:
            idx = index[col]
            return _cell_text(row[idx]) if idx < len(row) else ""

        records.append(
            {
                "row_no": i,
                "person_id": get("人员编号"),
                "item_raw": get("项目"),
                "result": get("结果"),
                "unit": canonical_unit(get("单位")),
                "exam_date": get("体检日期"),
            }
        )
        if len(records) >= MAX_ROWS:
            raise ExcelError(f"单次最多导入 {MAX_ROWS} 行，请拆分后上传")

    wb.close()
    if not records:
        raise ExcelError("未读取到任何数据行")
    return records


def parse_date(text: str) -> str | None:
    """接受 YYYY-MM-DD / YYYY/MM/DD / YYYYMMDD / datetime 字符串，输出 YYYY-MM-DD。"""
    t = (text or "").strip()
    if not t:
        return None
    fmts = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S")
    for f in fmts:
        try:
            return datetime.strptime(t, f).date().isoformat()
        except ValueError:
            continue
    return None


def parse_result_number(text: str) -> float | None:
    """结果数值化：'5.6' -> 5.6；非数字返回 None（布尔式 TRUE/FALSE 不算）。"""
    t = (text or "").strip()
    if not t:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def build_template_workbook(with_sample: bool = False) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "体检结果"
    ws.append(REQUIRED_HEADERS)

    samples = [
        ["EMP0001", "空腹血糖", 5.4, "mmol/L", "2026-09-20"],
        ["EMP0002", "收缩压", 148, "mmHg", "2026-09-20"],
        ["EMP0003", "BMI", 22.3, "kg/m2", "2026-09-21"],
        ["EMP0001", "空腹血糖", 5.4, "mmol/L", "2026-09-20"],   # 文件内重复
        ["EMP0004", "血红蛋白", 118, "g/L", "2026-09-21"],
        ["EMP0005", "总胆固醇", 5.9, "mmol/L", "2026-09-21"],
        ["EMP0006", "肌酐", 70, "umol/L", "2026-09-22"],
        ["EMP0007", "舒张压", 82, "mmHg", "2026-09-22"],
        ["EMP0008", "ALT", 33, "U/L", "2026-09-22"],
        ["EMP0009", "ALT", 35, "U/L", "2026-09-22"],
        ["EMP0010", "空腹血糖", 99, "mmol/L", "2026-09-22"],      # 平台硬校验
        ["EMP0500", "BMI", 21.0, "kg/m2", "2026-09-22"],          # 平台维护
        ["EMP0007", "收缩压", 120, "mmHg", "2026-09-22"],         # 演示限流（同批次第二次）
        ["EMP0007", "BMI", 23.0, "kg/m2", "2026-09-22"],
        ["EMP9999", "BMI", 20.5, "kg/m2", "2026-09-22"],          # 人员未建档
        ["EMP0011", "糖化血红蛋白", 6.2, "%", "2026-09-23"],       # 项目不在目录
        ["EMP0012", "空腹血糖", "", "mmol/L", "2026-09-23"],      # 结果为空
        ["EMP0013", "收缩压", 125, "kPa", "2026-09-23"],          # 单位不符
        ["EMP0014", "BMI", "偏高", "kg/m2", "2026-09-23"],        # 结果非数值
        ["EMP0015", "空腹血糖", 5.2, "mmol/L", "2026/9/23"],      # 日期格式可兼容
    ]
    if with_sample:
        for s in samples:
            ws.append(s)

    ws2 = wb.create_sheet("项目代码对照")
    ws2.append(["项目代码", "项目名称", "可接受写法", "标准单位", "参考区间下限", "参考区间上限"])
    from .catalog import ITEMS

    for it in ITEMS:
        ws2.append([
            it["code"], it["name"], " / ".join(it["aliases"]),
            "、".join(it["units"]), it["low"], it["high"],
        ])

    for col, width in zip("ABCDEF", [12, 18, 20, 12, 14, 14]):
        ws.column_dimensions[col].width = width
    for col, width in zip("ABCDEF", [12, 20, 28, 14, 14, 14]):
        ws2.column_dimensions[col].width = width
    return wb
