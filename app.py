#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""体检结果批量上传应用
工作人员导入 Excel（人员编号/项目/结果/单位/体检日期），后端校验后上送健康平台，
平台返回异常按行展示；支持任务进度、重复识别、异常导出与失败重试。
"""
import io
import os
import random
import re
import sqlite3
import threading
import uuid
from datetime import date, datetime

from flask import Flask, g, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
UPLOAD_DIR = os.path.join(DATA_DIR, "uploads")
DB_PATH = os.path.join(DATA_DIR, "app.db")
os.makedirs(UPLOAD_DIR, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # 16MB

ALL_HEADERS = ["人员编号", "项目", "结果", "单位", "体检日期"]
REQUIRED_HEADERS = ["人员编号", "项目", "结果", "体检日期"]

# 模拟健康平台的项目目录（真实环境由平台接口提供）
PLATFORM_ITEMS = {
    "身高", "体重", "血压", "心率", "体温", "血常规", "尿常规",
    "肝功能", "肾功能", "血糖", "血脂", "心电图", "B超", "胸片",
    "视力", "听力", "肺功能",
}

PERSON_ID_RE = re.compile(r"^[A-Za-z0-9\-]{4,32}$")
NOW = lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------- 数据库
def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS tasks (
            id          TEXT PRIMARY KEY,
            filename    TEXT NOT NULL,
            status      TEXT NOT NULL DEFAULT 'pending',  -- pending/processing/finished/failed
            total       INTEGER NOT NULL DEFAULT 0,
            processed   INTEGER NOT NULL DEFAULT 0,
            success     INTEGER NOT NULL DEFAULT 0,
            failed      INTEGER NOT NULL DEFAULT 0,
            duplicated  INTEGER NOT NULL DEFAULT 0,
            message     TEXT DEFAULT '',
            created_at  TEXT NOT NULL,
            finished_at TEXT
        );
        CREATE TABLE IF NOT EXISTS task_rows (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            task_id    TEXT NOT NULL,
            row_no     INTEGER NOT NULL,
            person_id  TEXT DEFAULT '',
            item       TEXT DEFAULT '',
            result     TEXT DEFAULT '',
            unit       TEXT DEFAULT '',
            exam_date  TEXT DEFAULT '',
            status     TEXT NOT NULL,   -- success/failed/duplicate
            stage      TEXT DEFAULT '', -- validation/platform/duplicate
            message    TEXT DEFAULT ''
        );
        CREATE INDEX IF NOT EXISTS idx_rows_task ON task_rows(task_id, status);
        -- 已成功上送的记录（人员编号+项目+体检日期唯一），用于跨任务历史去重
        CREATE TABLE IF NOT EXISTS uploaded (
            person_id TEXT NOT NULL,
            item      TEXT NOT NULL,
            exam_date TEXT NOT NULL,
            task_id   TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(person_id, item, exam_date)
        );
        """
    )
    conn.commit()
    conn.close()


# ---------------------------------------------------------------- 校验与平台
def parse_exam_date(value):
    """解析体检日期，返回 (iso字符串, 错误信息)"""
    if value is None or str(value).strip() == "":
        return None, "体检日期为空"
    if isinstance(value, datetime):
        d = value.date()
    elif isinstance(value, date):
        d = value
    else:
        s = str(value).strip()
        d = None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y年%m月%d日"):
            try:
                d = datetime.strptime(s, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            return None, "体检日期格式无法识别: %s" % s
    if d > date.today():
        return None, "体检日期晚于今天"
    if d.year < 2000:
        return None, "体检日期早于2000年，请核实"
    return d.isoformat(), None


def validate_row(person_id, item, result, unit):
    errors = []
    if not person_id:
        errors.append("人员编号为空")
    elif not PERSON_ID_RE.match(person_id):
        errors.append("人员编号须为4-32位字母/数字/横线")
    if not item:
        errors.append("项目为空")
    if not result:
        errors.append("结果为空")
    elif len(result) > 100:
        errors.append("结果超长(>100字符)")
    if unit and len(unit) > 20:
        errors.append("单位超长(>20字符)")
    return errors


def platform_upload(record):
    """模拟健康平台接口。真实环境替换为 HTTP 调用即可。
    返回 (是否成功, 平台消息)。模拟规则（便于演示按行异常）：
      - 人员编号含 404      -> 人员不存在
      - 项目不在平台目录    -> 项目未在平台目录中
      - 结果含 ERR          -> 结果格式不被平台接受
    """
    import time
    time.sleep(random.uniform(0.02, 0.05))  # 模拟网络耗时
    if "404" in record["person_id"]:
        return False, "平台返回：人员不存在"
    if record["item"] not in PLATFORM_ITEMS:
        return False, "平台返回：项目未在平台目录中"
    if "ERR" in record["result"].upper():
        return False, "平台返回：结果格式不被接受"
    return True, "OK"


# ---------------------------------------------------------------- 后台处理
def _save_row(conn, task_id, row_no, rec, status, stage, message):
    conn.execute(
        "INSERT INTO task_rows(task_id,row_no,person_id,item,result,unit,exam_date,status,stage,message)"
        " VALUES (?,?,?,?,?,?,?,?,?,?)",
        (task_id, row_no, rec["person_id"], rec["item"], rec["result"],
         rec["unit"], rec["exam_date"], status, stage, message),
    )


def _bump(conn, task_id, field):
    conn.execute(
        "UPDATE tasks SET processed=processed+1, %s=%s+1 WHERE id=?" % (field, field),
        (task_id,),
    )


def handle_one(conn, task_id, row_no, rec, seen):
    """处理单行：校验 -> 文件内去重 -> 历史去重 -> 上送平台。返回 (status, stage, message)"""
    errors = validate_row(rec["person_id"], rec["item"], rec["result"], rec["unit"])
    exam_date, date_err = parse_exam_date(rec["exam_date_raw"])
    if date_err:
        errors.append(date_err)
        # 保留原始输入，便于在明细与导出中排查
        rec["exam_date"] = str(rec["exam_date_raw"] or "").strip()
    else:
        rec["exam_date"] = exam_date
    if errors:
        return "failed", "validation", "；".join(errors)

    key = (rec["person_id"], rec["item"], rec["exam_date"])
    if key in seen:
        return "duplicate", "duplicate", "文件内重复（人员编号+项目+体检日期相同）"

    # 历史去重：占位插入，原子判断
    cur = conn.execute(
        "INSERT OR IGNORE INTO uploaded(person_id,item,exam_date,task_id,created_at)"
        " VALUES (?,?,?,?,?)",
        (rec["person_id"], rec["item"], rec["exam_date"], task_id, NOW()),
    )
    if cur.rowcount == 0:
        return "duplicate", "duplicate", "与历史已上传记录重复"

    ok, msg = platform_upload(rec)
    if ok:
        seen.add(key)
        return "success", "platform", "上送成功"
    # 平台退回：释放占位，便于重试
    conn.execute(
        "DELETE FROM uploaded WHERE person_id=? AND item=? AND exam_date=?",
        (rec["person_id"], rec["item"], rec["exam_date"]),
    )
    return "failed", "platform", msg


def process_task(task_id, filepath):
    conn = sqlite3.connect(DB_PATH)
    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        header_map, data_rows = None, []
        for i, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            if header_map is None:
                cells = ["" if c is None else str(c).strip() for c in raw]
                if set(REQUIRED_HEADERS).issubset(set(cells)):
                    header_map = {n: cells.index(n) for n in ALL_HEADERS if n in cells}
                continue
            if all(c is None or str(c).strip() == "" for c in raw):
                continue  # 跳过空行
            data_rows.append((i, raw))
        wb.close()

        if header_map is None:
            conn.execute(
                "UPDATE tasks SET status='failed', message=?, finished_at=? WHERE id=?",
                ("未找到表头行，需包含：%s" % "、".join(REQUIRED_HEADERS), NOW(), task_id),
            )
            conn.commit()
            return

        conn.execute("UPDATE tasks SET status='processing', total=? WHERE id=?",
                     (len(data_rows), task_id))
        conn.commit()

        seen = set()
        for row_no, raw in data_rows:
            def col(name):
                idx = header_map.get(name)
                return raw[idx] if idx is not None and idx < len(raw) else None

            rec = {
                "person_id": str(col("人员编号") or "").strip(),
                "item": str(col("项目") or "").strip(),
                "result": str(col("结果") or "").strip(),
                "unit": str(col("单位") or "").strip(),
                "exam_date": "",
                "exam_date_raw": col("体检日期"),
            }
            status, stage, msg = handle_one(conn, task_id, row_no, rec, seen)
            _save_row(conn, task_id, row_no, rec, status, stage, msg)
            _bump(conn, task_id,
                  {"success": "success", "failed": "failed", "duplicate": "duplicated"}[status])
            if row_no % 20 == 0:
                conn.commit()

        conn.execute("UPDATE tasks SET status='finished', finished_at=? WHERE id=?",
                     (NOW(), task_id))
        conn.commit()
    except Exception as exc:  # 文件损坏等
        conn.execute("UPDATE tasks SET status='failed', message=?, finished_at=? WHERE id=?",
                     ("处理失败：%s" % exc, NOW(), task_id))
        conn.commit()
    finally:
        conn.close()


def retry_task(task_id):
    """仅重试被平台退回的行（stage=platform），校验失败与重复行不重试"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("UPDATE tasks SET status='processing', message='' WHERE id=?", (task_id,))
        conn.commit()
        rows = conn.execute(
            "SELECT * FROM task_rows WHERE task_id=? AND status='failed' AND stage='platform'"
            " ORDER BY row_no", (task_id,)).fetchall()
        seen = set()
        for r in rows:
            rec = {"person_id": r["person_id"], "item": r["item"], "result": r["result"],
                   "unit": r["unit"], "exam_date": r["exam_date"], "exam_date_raw": r["exam_date"]}
            status, stage, msg = handle_one(conn, task_id, r["row_no"], rec, seen)
            conn.execute("UPDATE task_rows SET status=?, stage=?, message=? WHERE id=?",
                         (status, stage, msg, r["id"]))
            conn.commit()
        # 重算任务计数
        conn.execute(
            """UPDATE tasks SET
                 success    =(SELECT COUNT(*) FROM task_rows WHERE task_id=? AND status='success'),
                 failed     =(SELECT COUNT(*) FROM task_rows WHERE task_id=? AND status='failed'),
                 duplicated =(SELECT COUNT(*) FROM task_rows WHERE task_id=? AND status='duplicate'),
                 status='finished', finished_at=?
               WHERE id=?""",
            (task_id, task_id, task_id, NOW(), task_id))
        conn.commit()
    finally:
        conn.close()


# ---------------------------------------------------------------- 接口
@app.route("/")
def index():
    return render_template("index.html", platform_items=sorted(PLATFORM_ITEMS))


@app.post("/api/upload")
def upload():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="请选择要上传的文件"), 400
    if not f.filename.lower().endswith((".xlsx", ".xlsm")):
        return jsonify(error="仅支持 .xlsx 格式（.xls 请另存为 .xlsx）"), 400
    task_id = uuid.uuid4().hex[:12]
    path = os.path.join(UPLOAD_DIR, task_id + ".xlsx")
    f.save(path)
    db = get_db()
    db.execute("INSERT INTO tasks(id,filename,status,created_at) VALUES (?,?,?,?)",
               (task_id, f.filename, "pending", NOW()))
    db.commit()
    threading.Thread(target=process_task, args=(task_id, path), daemon=True).start()
    return jsonify(task_id=task_id)


@app.get("/api/tasks")
def tasks():
    rows = get_db().execute(
        "SELECT * FROM tasks ORDER BY created_at DESC, id DESC LIMIT 50").fetchall()
    return jsonify([dict(r) for r in rows])


@app.get("/api/tasks/<task_id>/rows")
def task_rows(task_id):
    status = request.args.get("status", "all")
    page = max(1, int(request.args.get("page", 1)))
    size = 50
    where, args = "task_id=?", [task_id]
    if status in ("success", "failed", "duplicate"):
        where += " AND status=?"
        args.append(status)
    db = get_db()
    total = db.execute("SELECT COUNT(*) c FROM task_rows WHERE " + where, args).fetchone()["c"]
    rows = db.execute(
        "SELECT * FROM task_rows WHERE " + where + " ORDER BY row_no LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    return jsonify(total=total, page=page, size=size, rows=[dict(r) for r in rows])


@app.get("/api/tasks/<task_id>/export")
def export(task_id):
    db = get_db()
    task = db.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
    if not task:
        return jsonify(error="任务不存在"), 404
    rows = db.execute(
        "SELECT * FROM task_rows WHERE task_id=? AND status IN ('failed','duplicate')"
        " ORDER BY row_no", (task_id,)).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "异常明细"
    ws.append(["行号", "人员编号", "项目", "结果", "单位", "体检日期", "异常类型", "异常原因"])
    stage_text = {"validation": "校验失败", "platform": "平台退回", "duplicate": "重复"}
    for r in rows:
        ws.append([r["row_no"], r["person_id"], r["item"], r["result"], r["unit"],
                   r["exam_date"], stage_text.get(r["stage"], r["stage"]), r["message"]])
    for col, width in zip("ABCDEFGH", (8, 14, 12, 16, 10, 14, 12, 44)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = "异常明细_%s.xlsx" % task_id
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/tasks/<task_id>/retry")
def retry(task_id):
    db = get_db()
    task = db.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
    if not task:
        return jsonify(error="任务不存在"), 404
    if task["status"] in ("pending", "processing"):
        return jsonify(error="任务正在处理中，请稍后再试"), 409
    n = db.execute(
        "SELECT COUNT(*) c FROM task_rows WHERE task_id=? AND status='failed' AND stage='platform'",
        (task_id,)).fetchone()["c"]
    if n == 0:
        return jsonify(error="没有可重试的平台退回行"), 400
    threading.Thread(target=retry_task, args=(task_id,), daemon=True).start()
    return jsonify(ok=True, retrying=n)


@app.get("/api/template")
def template():
    wb = Workbook()
    ws = wb.active
    ws.title = "体检结果"
    ws.append(ALL_HEADERS)
    ws.append(["P10001", "血压", "120/80", "mmHg", "2026-09-20"])
    ws.append(["P10001", "心率", "72", "次/分", "2026-09-20"])
    ws.append(["P10002", "血糖", "5.6", "mmol/L", "2026-09-21"])
    for col, width in zip("ABCDE", (14, 12, 14, 10, 14)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="体检结果导入模板.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, threaded=True)
