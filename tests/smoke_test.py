"""端到端冒烟测试（不依赖外部服务，使用 Flask test client）。

运行：
  HC_PLATFORM_FAIL_RATE=0 .venv/bin/python tests/smoke_test.py
"""
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# 隔离数据目录 & 关闭随机故障，保证结果可重复
TMP = tempfile.mkdtemp(prefix="hc_test_")
os.environ["HC_DATA_DIR"] = TMP
os.environ.setdefault("HC_PLATFORM_FAIL_RATE", "0")
os.environ.setdefault("HC_PLATFORM_LATENCY_MS", "0")

from openpyxl import load_workbook

from app import create_app
from app import db, statuses as st
from app.platform import client as platform_client
platform_client.rate_limit_per_sec = 100_000  # 测试环境不限速

SAMPLE = ROOT / "samples" / "体检结果_示例.xlsx"
app = create_app()
c = app.test_client()

passed, failed = [], []


def check(name, cond, detail=""):
    (passed if cond else failed).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name} {detail}")
    assert cond, f"{name} {detail}"


def wait_done(task_id, timeout=30):
    for _ in range(timeout * 10):
        t = db.get_task(task_id)
        if t["status"] not in st.ACTIVE_TASK_STATUS:
            return t
        time.sleep(0.1)
    raise AssertionError(f"task {task_id} timeout")


print("== 1. 上传示例文件 ==")
with open(SAMPLE, "rb") as f:
    rv = c.post("/api/tasks", data={"file": (f, SAMPLE.name)}, content_type="multipart/form-data")
check("上传返回 201", rv.status_code == 201, str(rv.status_code))
task_id = rv.get_json()["task_id"]
t = wait_done(task_id)
print(f"     任务 {task_id} 最终状态：{t['status']}")
check("总行为 35", t["total_count"] == 35, f"total={t['total_count']}")
check("校验失败 9 行", t["invalid_count"] == 9, f"invalid={t['invalid_count']}")
check("文件内重复 1 行", t["duplicate_count"] == 1, f"dup={t['duplicate_count']}")
check("平台异常 1 行（查无此人 P9999）", t["platform_error_count"] == 1,
      f"perr={t['platform_error_count']}")
check("成功 24 行", t["success_count"] == 24, f"success={t['success_count']}")
check("任务状态=完成（含异常）", t["status"] == st.TASK_COMPLETED_WITH_ERRORS, t["status"])

print("== 2. 警告行（超参考范围但上送成功） ==")
warn_rows = [r for r in db.list_rows(task_id, st.SUCCESS, limit=1000) if r["warning"]]
check("至少 5 条超参考范围警告", len(warn_rows) >= 5, f"warn={len(warn_rows)}")

print("== 3. 平台异常逐行可查询 ==")
rv = c.get(f"/api/tasks/{task_id}/rows?status=PLATFORM_ERROR")
data = rv.get_json()
check("平台异常接口 total=1", data["total"] == 1)
check("错误码=PERSON_NOT_FOUND", data["rows"][0]["error_code"] == "PERSON_NOT_FOUND",
      data["rows"][0]["error_code"])

print("== 4. 异常导出 Excel ==")
rv = c.get(f"/api/tasks/{task_id}/errors.xlsx")
check("导出状态 200", rv.status_code == 200)
xlsx_path = Path(TMP) / "errors.xlsx"
xlsx_path.write_bytes(rv.data)
wb = load_workbook(xlsx_path)
check("含汇总页", "汇总" in wb.sheetnames)
check("含校验失败页", "校验失败" in wb.sheetnames)
check("含平台异常页", "平台异常" in wb.sheetnames)
check("含重复数据页", "重复数据" in wb.sheetnames)
ws_inv = wb["校验失败"]
check("校验失败页 9 条数据", sum(1 for row in ws_inv.iter_rows(min_row=2, values_only=True)
                          if row[0] is not None and isinstance(row[0], int)) == 9)
ws_dup = wb["重复数据"]
check("重复数据页 1 条数据", sum(1 for row in ws_dup.iter_rows(min_row=2, values_only=True)
                          if row[0] is not None and isinstance(row[0], int)) == 1)

print("== 5. 跨文件去重 ==")
# 第二份文件：P1011/WBC/2026-09-08 已在第一份文件成功上送并占用指纹
from openpyxl import Workbook
wb2 = Workbook()
ws2 = wb2.active
ws2.append(["人员编号", "项目", "结果", "单位", "体检日期"])
ws2.append(["P1011", "WBC", 7.1, "10^9/L", "2026-09-08"])  # 跨文件重复
ws2.append(["P1020", "GLU", 4.8, "mmol/L", "2026-09-08"])  # 新增正常行
f2 = Path(TMP) / "batch2.xlsx"
wb2.save(f2)
with open(f2, "rb") as fh:
    rv = c.post("/api/tasks", data={"file": (fh, "batch2.xlsx")}, content_type="multipart/form-data")
tid2 = rv.get_json()["task_id"]
t2 = wait_done(tid2)
check("批次2 重复 1 行", t2["duplicate_count"] == 1, f"dup={t2['duplicate_count']}")
check("批次2 成功 1 行", t2["success_count"] == 1, f"success={t2['success_count']}")
dup_row = [r for r in db.list_rows(tid2, st.DUPLICATE, limit=10)][0]
check("重复错误码=跨文件", dup_row["error_code"] == "DUPLICATE_CROSS_FILE", dup_row["error_code"])
check("重复原因含源文件", "体检结果_示例.xlsx" in dup_row["error_message"], dup_row["error_message"])

print("== 6. 平台幂等（同指纹直接重上送应被平台 409 拦截） ==")
from app.processor import manager
# 构造一个“假待上送行”验证平台侧幂等记忆
rec = {"fingerprint": "test-fp-xyz", "person_id": "P1001", "item_code": "GLU",
       "exam_date": "2026-09-01", "result": "5.2"}
r1 = platform_client.upload(rec)
check("首次上送返回记录号", bool(r1["record_id"]))
try:
    platform_client.upload(rec)
    check("重复上送抛异常", False)
except Exception as e:
    check("重复上送被幂等拦截", e.code == "DUPLICATE_SUBMISSION", e.code)

print("== 7. 错误文件被拒绝（缺表头） ==")
bad = Path(TMP) / "bad.xlsx"
wb3 = Workbook()
wb3.active.append(["foo", "bar"])
wb3.active.append([1, 2])
wb3.save(bad)
with open(bad, "rb") as fh:
    rv = c.post("/api/tasks", data={"file": (fh, "bad.xlsx")}, content_type="multipart/form-data")
check("缺表头返回 400", rv.status_code == 400, str(rv.status_code))
check("错误信息提示表头", "表头" in rv.get_json()["error"])

print("== 8. 模板下载 ==")
rv = c.get("/api/template")
check("模板 200", rv.status_code == 200)
tpl = Path(TMP) / "tpl.xlsx"
tpl.write_bytes(rv.data)
wb4 = load_workbook(tpl)
check("模板含 3 个页签", set(wb4.sheetnames) == {"体检结果", "填写说明", "项目目录"}, str(wb4.sheetnames))

print("== 9. 取消 & 继续上送 ==")
# 调大平台延迟，使任务处于“上送中”时可以被取消
platform_client.rate_limit_per_sec = 10_000
platform_client.latency_ms = 60
wb5 = Workbook()
ws5 = wb5.active
ws5.append(["人员编号", "项目", "结果", "单位", "体检日期"])
for i in range(12):
    ws5.append([f"P10{30+i}", "GLU", 5.0 + i * 0.01, "mmol/L", "2026-09-09"])
f5 = Path(TMP) / "batch3.xlsx"
wb5.save(f5)
with open(f5, "rb") as fh:
    rv = c.post("/api/tasks", data={"file": (fh, "batch3.xlsx")}, content_type="multipart/form-data")
tid3 = rv.get_json()["task_id"]
time.sleep(0.25)
rv = c.post(f"/api/tasks/{tid3}/cancel")
check("取消请求 200", rv.status_code == 200, str(rv.status_code))
for _ in range(50):
    t3 = db.get_task(tid3)
    if t3["status"] not in st.ACTIVE_TASK_STATUS:
        break
    time.sleep(0.1)
check("任务最终为已取消", t3["status"] == st.TASK_CANCELLED, t3["status"])
check("取消后仍有待上送行", t3["pending_count"] > 0, f"pending={t3['pending_count']}")
platform_client.latency_ms = 0
rv = c.post(f"/api/tasks/{tid3}/resume")
check("继续上送 200", rv.status_code == 200, str(rv.status_code))
t3 = wait_done(tid3)
check("继续后全部成功", t3["success_count"] == 12, f"success={t3['success_count']}")

print("\n========== 结果 ==========")
print(f"通过 {len(passed)} 项，失败 {len(failed)} 项")
if failed:
    print("失败项：", failed)
    sys.exit(1)
print("全部通过 ✔")
